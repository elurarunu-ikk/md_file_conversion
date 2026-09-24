#!/usr/bin/env python3
"""Convert common document formats into Markdown files.

This script scans an input folder for .docx, .xlsx, .html/.htm, and .pdf files,
converts each file to Markdown, and writes the results to an output folder.
It includes a simple Tkinter GUI for selecting folders and a resilient
error-handling workflow to keep processing even when a file fails.

Requirements:
    pip install python-docx openpyxl markdownify beautifulsoup4 pypdf
"""

from __future__ import annotations

import datetime as dt
import logging
import os
import re
import threading
import traceback
from pathlib import Path
from typing import Callable, List, Optional

import tkinter as tk
from tkinter import filedialog, messagebox, ttk

TARGET_EXTENSIONS = {".docx", ".xlsx", ".html", ".htm", ".pdf"}

# Office lock/temp files (~$data.xlsx, .~lock.data.xlsx#) and hidden files.
SKIP_PREFIXES = ("~$", ".~", ".")

try:
    from markdownify import markdownify
except ImportError:  # pragma: no cover
    markdownify = None

try:
    from bs4 import BeautifulSoup, Comment
except ImportError:  # pragma: no cover
    BeautifulSoup = None

try:
    from docx import Document
except ImportError:  # pragma: no cover
    Document = None

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover
    load_workbook = None

try:
    from pypdf import PdfReader
    from pypdf.errors import DependencyError
except ImportError:  # pragma: no cover
    PdfReader = None
    DependencyError = None


def configure_logger(log_path: Path) -> logging.Logger:
    logger = logging.getLogger("md_converter")
    logger.setLevel(logging.INFO)

    for handler in list(logger.handlers):
        handler.close()
    logger.handlers.clear()

    formatter = logging.Formatter("%(asctime)s | %(levelname)s | %(message)s")
    file_handler = logging.FileHandler(log_path, encoding="utf-8")
    file_handler.setFormatter(formatter)
    logger.addHandler(file_handler)
    return logger


def clean_text(value: str) -> str:
    """Normalize extracted text for Markdown output."""
    if value is None:
        return ""
    cleaned = value.replace("\r\n", "\n").replace("\r", "\n")
    cleaned = re.sub(r"\n{3,}", "\n\n", cleaned)
    cleaned = re.sub(r"[ \t]+\n", "\n", cleaned)
    return cleaned.strip()


def escape_markdown_cell(value: object) -> str:
    text = "" if value is None else str(value)
    return text.replace("|", "\\|").replace("\n", "<br>").strip()


def table_to_markdown(rows: List[List[object]]) -> str:
    """Convert a 2D table into Markdown pipe syntax."""
    if not rows:
        return ""

    width = max(len(row) for row in rows)
    normalized = [list(row) + [""] * (width - len(row)) for row in rows]
    escaped = [[escape_markdown_cell(cell) for cell in row] for row in normalized]

    header = escaped[0]
    separator = ["---"] * width
    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join(separator) + " |",
    ]

    for row in escaped[1:]:
        lines.append("| " + " | ".join(row) + " |")

    return "\n".join(lines)


def docx_to_markdown(file_path: Path) -> str:
    """Convert a DOCX file to Markdown using python-docx."""
    if Document is None:
        raise RuntimeError("python-docx is not installed. Install with: pip install python-docx")

    document = Document(str(file_path))
    blocks: List[str] = []
    list_items: List[str] = []

    def flush_list() -> None:
        if list_items:
            blocks.append("\n".join(list_items))
            list_items.clear()

    # iter_inner_content() yields paragraphs and tables in document order.
    for item in document.iter_inner_content():
        if hasattr(item, "rows"):
            flush_list()
            rows = [[cell.text.strip() for cell in row.cells] for row in item.rows]
            markdown_table = table_to_markdown(rows)
            if markdown_table:
                blocks.append(markdown_table)
            continue

        text = clean_text(item.text)
        if not text:
            continue

        style_name = (item.style.name if item.style is not None else "") or ""
        heading = re.match(r"Heading (\d)", style_name)
        is_numbered_list = style_name.startswith("List Number")
        is_list = style_name.startswith("List") or (item._p.pPr is not None and item._p.pPr.numPr is not None)

        if is_list:
            list_items.append(("1. " if is_numbered_list else "- ") + text)
            continue

        flush_list()
        if style_name == "Title":
            blocks.append(f"# {text}")
        elif heading:
            level = min(int(heading.group(1)), 6)
            blocks.append(f"{'#' * level} {text}")
        else:
            blocks.append(text)

    flush_list()
    return "\n\n".join(blocks) if blocks else "# DOCX document\n\nNo readable content found."


def format_cell(value: object) -> str:
    """Render a spreadsheet value: dates as 2026-01-15, no .0 on whole numbers."""
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        if value.time() == dt.time(0, 0):
            return value.strftime("%Y-%m-%d")
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, dt.date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, dt.time):
        return value.strftime("%H:%M:%S")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def xlsx_to_markdown(file_path: Path) -> str:
    """Convert an XLSX workbook to Markdown, one sheet at a time."""
    if load_workbook is None:
        raise RuntimeError("openpyxl is not installed. Install with: pip install openpyxl")

    workbook_parts: List[str] = []
    workbook = load_workbook(file_path, read_only=True, data_only=True)

    try:
        for sheet in workbook.worksheets:
            # Some writers store a wrong <dimension>; without this, read-only
            # mode can stop early (e.g. only A1).
            sheet.reset_dimensions()

            rows = [
                [format_cell(value) for value in row]
                for row in sheet.iter_rows(values_only=True)
            ]
            rows = [row for row in rows if any(cell.strip() for cell in row)]
            if not rows:
                continue

            width = max(
                max((i + 1 for i, cell in enumerate(row) if cell.strip()), default=0)
                for row in rows
            )
            rows = [row[:width] for row in rows]

            sheet_markdown = table_to_markdown(rows)
            workbook_parts.append(f"## {sheet.title}\n\n{sheet_markdown}")
    finally:
        workbook.close()

    if not workbook_parts:
        return "# Spreadsheet\n\nNo worksheet data found."

    return "\n\n".join(workbook_parts)


def html_to_markdown(file_path: Path) -> str:
    """Convert HTML content to Markdown using markdownify."""
    if markdownify is None:
        raise RuntimeError("markdownify is not installed. Install with: pip install markdownify")

    if BeautifulSoup is None:
        raise RuntimeError("beautifulsoup4 is not installed. Install with: pip install beautifulsoup4")

    html_content = file_path.read_text(encoding="utf-8", errors="replace")
    soup = BeautifulSoup(html_content, "html.parser")

    title = soup.title.get_text(" ", strip=True) if soup.title else ""

    # markdownify's strip= keeps the inner text of stripped tags, so remove
    # these elements (with their contents) before converting.
    for tag in soup.find_all(["head", "title", "script", "style", "noscript", "template"]):
        tag.decompose()
    for comment in soup.find_all(string=lambda node: isinstance(node, Comment)):
        comment.extract()

    body = soup.body or soup
    markdown_text = clean_text(markdownify(str(body), heading_style="ATX"))

    if title and body.find("h1") is None:
        markdown_text = f"# {title}\n\n{markdown_text}".strip()

    return markdown_text or "# HTML document\n\nNo readable content found."


def pdf_to_markdown(file_path: Path) -> str:
    """Convert a PDF file to Markdown by extracting text page by page."""
    if PdfReader is None:
        raise RuntimeError("pypdf is not installed. Install with: pip install pypdf")

    reader = PdfReader(str(file_path))
    if reader.is_encrypted:
        # Many PDFs are "encrypted" with an empty user password (print/copy
        # restrictions only) and open fine with "".
        try:
            decrypted = reader.decrypt("")
        except DependencyError as exc:
            raise RuntimeError(
                "PDF uses AES encryption; install with: pip install cryptography"
            ) from exc
        except Exception as exc:  # noqa: BLE001
            raise RuntimeError("PDF is password-protected") from exc
        if not decrypted:
            raise RuntimeError("PDF is password-protected")

    pages: List[str] = []

    for page in reader.pages:
        text = page.extract_text() or ""
        cleaned = clean_text(text)
        if cleaned:
            pages.append(cleaned)

    if not pages:
        return "# PDF document\n\nNo extractable text found (the PDF may be scanned images)."

    return "\n\n---\n\n".join(pages)


def convert_file_to_markdown(file_path: Path) -> str:
    suffix = file_path.suffix.lower()

    if suffix == ".docx":
        return docx_to_markdown(file_path)
    if suffix == ".xlsx":
        return xlsx_to_markdown(file_path)
    if suffix in {".html", ".htm"}:
        return html_to_markdown(file_path)
    if suffix == ".pdf":
        return pdf_to_markdown(file_path)

    raise ValueError(f"Unsupported file type: {suffix}")


def build_output_path(input_dir: Path, output_dir: Path, source_file: Path, used: set) -> Path:
    """Mirror the source's subfolder under output_dir.

    The name is deterministic so re-running overwrites instead of duplicating.
    If two sources in one folder share a stem (report.docx, report.pdf), the
    later one gets its extension appended (report_pdf.md).
    """
    target_dir = output_dir / source_file.parent.relative_to(input_dir)
    candidate = target_dir / f"{source_file.stem}.md"
    if candidate in used:
        candidate = target_dir / f"{source_file.stem}_{source_file.suffix.lstrip('.').lower()}.md"
    used.add(candidate)
    return candidate


def find_target_files(input_dir: Path, exclude_dir: Path | None = None) -> List[Path]:
    matches: List[Path] = []
    for root, dirs, files in os.walk(input_dir):
        # Skip hidden folders, and don't scan the output folder if it lives
        # inside the input folder.
        dirs[:] = [
            d
            for d in dirs
            if not d.startswith(".") and (exclude_dir is None or (Path(root) / d).resolve() != exclude_dir)
        ]
        for file_name in files:
            if file_name.startswith(SKIP_PREFIXES):
                continue
            path = Path(root) / file_name
            if path.suffix.lower() in TARGET_EXTENSIONS:
                matches.append(path)
    return sorted(matches)


def convert_folder(
    input_dir: Path,
    output_dir: Path,
    logger: logging.Logger,
    progress: Optional[Callable[[str], None]] = None,
) -> dict:
    """Convert every supported file under input_dir into output_dir.

    If given, progress is called once per file with a line such as
    "[3/10] OK     sub\\report.docx" or "[4/10] FAILED a.pdf -> reason".
    """
    input_dir = input_dir.expanduser().resolve()
    output_dir = output_dir.expanduser().resolve()
    output_dir.mkdir(parents=True, exist_ok=True)

    files = find_target_files(input_dir, exclude_dir=output_dir)
    if not files:
        raise FileNotFoundError(
            f"No supported files found in '{input_dir}'. Accepted extensions: {sorted(TARGET_EXTENSIONS)}"
        )

    stats = {"processed": 0, "errors": 0}
    used_paths: set = set()

    total = len(files)

    for index, file_path in enumerate(files, start=1):
        relative = file_path.relative_to(input_dir)
        try:
            markdown_text = convert_file_to_markdown(file_path)
            output_path = build_output_path(input_dir, output_dir, file_path, used_paths)
            output_path.parent.mkdir(parents=True, exist_ok=True)
            output_path.write_text(markdown_text, encoding="utf-8")
            stats["processed"] += 1
            logger.info("Converted %s -> %s", relative, output_path.relative_to(output_dir))
            message = f"[{index}/{total}] OK     {relative}"
        except Exception as exc:  # noqa: BLE001
            stats["errors"] += 1
            logger.error("Failed to convert %s: %s", relative, exc)
            logger.error(traceback.format_exc())
            message = f"[{index}/{total}] FAILED {relative} -> {exc}"

        if progress is not None:
            progress(message)

    return stats


class FileToMarkdownApp(tk.Tk):
    """Tkinter GUI for selecting folders and starting conversion."""

    def __init__(self) -> None:
        super().__init__()
        self.title("Document to Markdown Converter")
        self.geometry("760x480")
        self.minsize(650, 400)

        self.input_var = tk.StringVar()
        self.output_var = tk.StringVar()
        self.status_var = tk.StringVar(value="Ready")

        self._build_ui()

    def _build_ui(self) -> None:
        main = ttk.Frame(self, padding=18)
        main.pack(fill="both", expand=True)

        title = ttk.Label(main, text="Convert files to Markdown", font=("Segoe UI", 14, "bold"))
        title.grid(row=0, column=0, columnspan=3, sticky="w", pady=(0, 16))

        ttk.Label(main, text="Input folder:", font=("Segoe UI", 10, "bold")).grid(
            row=1, column=0, sticky="w", padx=(0, 10), pady=(0, 8)
        )
        ttk.Entry(main, textvariable=self.input_var, width=62).grid(row=1, column=1, sticky="ew")
        ttk.Button(main, text="Browse", command=self.choose_input_folder, width=12).grid(
            row=1, column=2, padx=(10, 0), sticky="ew"
        )

        ttk.Label(main, text="Output folder:", font=("Segoe UI", 10, "bold")).grid(
            row=2, column=0, sticky="w", padx=(0, 10), pady=(10, 8)
        )
        ttk.Entry(main, textvariable=self.output_var, width=62).grid(row=2, column=1, sticky="ew")
        ttk.Button(main, text="Browse", command=self.choose_output_folder, width=12).grid(
            row=2, column=2, padx=(10, 0), sticky="ew"
        )

        self.run_button = ttk.Button(main, text="Convert files", command=self.start_conversion)
        self.run_button.grid(row=3, column=0, columnspan=3, pady=(18, 8), sticky="ew")

        self.progress_bar = ttk.Progressbar(main, mode="determinate", maximum=1, value=0)
        self.progress_bar.grid(row=4, column=0, columnspan=3, sticky="ew", pady=(0, 12))

        ttk.Label(main, text="Status", font=("Segoe UI", 10, "bold")).grid(
            row=5, column=0, columnspan=3, sticky="w", pady=(0, 4)
        )

        self.status_label = ttk.Label(main, textvariable=self.status_var, foreground="#1f5f8b")
        self.status_label.grid(row=6, column=0, columnspan=3, sticky="w", pady=(0, 8))

        log_frame = ttk.Frame(main)
        log_frame.grid(row=7, column=0, columnspan=3, sticky="nsew")
        log_frame.columnconfigure(0, weight=1)
        log_frame.rowconfigure(0, weight=1)

        self.log_box = tk.Text(log_frame, height=10, wrap="word", bg="#f8f8f8", relief="solid", borderwidth=1)
        self.log_box.grid(row=0, column=0, sticky="nsew")
        log_scrollbar = ttk.Scrollbar(log_frame, orient="vertical", command=self.log_box.yview)
        log_scrollbar.grid(row=0, column=1, sticky="ns")
        self.log_box.configure(yscrollcommand=log_scrollbar.set)
        self.log_box.insert("end", "Ready. Select an input folder and output folder to begin.\n")
        self.log_box.configure(state="disabled")

        main.columnconfigure(1, weight=1)
        main.rowconfigure(7, weight=1)

    def append_log(self, message: str) -> None:
        self.log_box.configure(state="normal")
        self.log_box.insert("end", message + "\n")
        self.log_box.see("end")
        self.log_box.configure(state="disabled")

    def choose_input_folder(self) -> None:
        folder = filedialog.askdirectory(title="Select input folder")
        if folder:
            self.input_var.set(folder)
            self.append_log(f"Input folder selected: {folder}")
            if not self.output_var.get().strip():
                default_output = str(Path(folder) / "markdown_output")
                self.output_var.set(default_output)
                self.append_log(f"Output folder set to: {default_output}")

    def choose_output_folder(self) -> None:
        folder = filedialog.askdirectory(title="Select output folder")
        if folder:
            self.output_var.set(folder)
            self.append_log(f"Output folder selected: {folder}")

    def set_status(self, message: str) -> None:
        self.status_var.set(message)

    def on_progress(self, message: str) -> None:
        """Runs on the Tk thread; message looks like "[i/N] OK     path"."""
        match = re.match(r"\[(\d+)/(\d+)\]", message)
        if match:
            done, total = int(match.group(1)), int(match.group(2))
            self.progress_bar.configure(maximum=total, value=done)
            self.set_status(f"Converting... {done} of {total}")
        self.append_log(message)

    def start_conversion(self) -> None:
        input_dir = self.input_var.get().strip()
        output_dir = self.output_var.get().strip()

        if not input_dir:
            messagebox.showerror("Input required", "Please select an input folder.")
            return

        if not output_dir:
            messagebox.showerror("Output required", "Please select an output folder.")
            return

        input_path = Path(input_dir).expanduser()
        output_path = Path(output_dir).expanduser()

        if not input_path.exists() or not input_path.is_dir():
            messagebox.showerror("Invalid input", f"The input folder does not exist: {input_path}")
            return

        output_path.mkdir(parents=True, exist_ok=True)
        self.set_status("Scanning files and converting...")
        self.append_log(f"Processing folder: {input_path}")
        self.run_button.config(state="disabled")
        self.progress_bar.configure(maximum=1, value=0)

        thread = threading.Thread(
            target=self._conversion_worker,
            args=(input_path, output_path),
            daemon=True,
        )
        thread.start()

    def _conversion_worker(self, input_dir: Path, output_dir: Path) -> None:
        try:
            log_path = output_dir / "conversion.log"
            logger = configure_logger(log_path)
            stats = convert_folder(
                input_dir,
                output_dir,
                logger,
                progress=lambda message: self.after(0, self.on_progress, message),
            )
            completed_summary = (
                f"Completed: {stats['processed']} files converted, "
                f"{stats['errors']} files failed. Details are logged in {log_path}."
            )
            self.after(0, self.set_status, completed_summary)
            self.after(0, self.append_log, completed_summary)
            self.after(0, messagebox.showinfo, "Conversion complete", completed_summary)
        except Exception as exc:  # noqa: BLE001
            error_message = f"Conversion failed: {exc}"
            self.after(0, self.set_status, error_message)
            self.after(0, self.append_log, error_message)
            self.after(0, messagebox.showerror, "Conversion failed", error_message)
        finally:
            # Release conversion.log so the output folder can be moved/deleted
            # while the app stays open.
            for handler in list(logging.getLogger("md_converter").handlers):
                handler.close()
            self.after(0, lambda: self.run_button.config(state="normal"))


if __name__ == "__main__":
    app = FileToMarkdownApp()
    app.mainloop()
